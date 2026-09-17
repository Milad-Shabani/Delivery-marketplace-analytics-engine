import sys, os, json
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.formatting.rule import ColorScaleRule

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PROC = os.path.join(ROOT, "data", "processed")
RAW = os.path.join(ROOT, "data", "raw")
OUT = os.path.join(ROOT, "outputs", "excel", "NOVAFOOD_Business_Analytics.xlsx")

FONT = "Arial"
NAVY = "1F2A44"
TEAL = "0F7C82"
LIGHT = "EDF2F7"
WHITE = "FFFFFF"

TITLE_FONT = Font(name=FONT, size=16, bold=True, color=WHITE)
HEADER_FONT = Font(name=FONT, size=10, bold=True, color=WHITE)
LABEL_FONT = Font(name=FONT, size=10, bold=True, color=NAVY)
BODY_FONT = Font(name=FONT, size=10, color="1A1A1A")
NOTE_FONT = Font(name=FONT, size=9, italic=True, color="666666")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
BAND_FILL = PatternFill("solid", fgColor=LIGHT)
TITLE_FILL = PatternFill("solid", fgColor=NAVY)
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def sheet_title(ws: Worksheet, text: str, subtitle: str = "", width=12):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    c = ws.cell(row=1, column=1, value=text)
    c.font = TITLE_FONT
    c.fill = TITLE_FILL
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 28
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
        c2 = ws.cell(row=2, column=1, value=subtitle)
        c2.font = NOTE_FONT
        ws.row_dimensions[2].height = 16


def write_df(ws: Worksheet, df: pd.DataFrame, start_row: int, start_col=1, number_formats: dict | None = None,
             pct_cols: list | None = None, currency_cols: list | None = None, freeze=True) -> int:
    number_formats = number_formats or {}
    pct_cols = pct_cols or []
    currency_cols = currency_cols or []
    for j, col in enumerate(df.columns):
        cell = ws.cell(row=start_row, column=start_col + j, value=str(col).replace("_", " ").title())
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for i, (_, row) in enumerate(df.iterrows()):
        r = start_row + 1 + i
        for j, col in enumerate(df.columns):
            val = row[col]
            if isinstance(val, (pd.Timestamp,)):
                val = val.strftime("%Y-%m")
            if isinstance(val, (np.integer,)):
                val = int(val)
            if isinstance(val, (np.floating,)):
                val = float(val)
            if pd.isna(val):
                val = None
            cell = ws.cell(row=r, column=start_col + j, value=val)
            cell.font = BODY_FONT
            cell.border = BORDER
            if i % 2 == 1:
                cell.fill = BAND_FILL
            if col in pct_cols:
                cell.number_format = "0.0%"
            elif col in currency_cols:
                cell.number_format = "$#,##0.00"
            elif col in number_formats:
                cell.number_format = number_formats[col]
    for j, col in enumerate(df.columns):
        letter = get_column_letter(start_col + j)
        maxlen = max(len(str(col)), *(len(str(v)) for v in df[col].astype(str).values[:200])) if len(df) else len(str(col))
        ws.column_dimensions[letter].width = min(max(12, maxlen + 2), 34)
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=start_col).coordinate
    return start_row + 1 + len(df)


def add_line_chart(ws, title, data_ref, cats_ref, anchor, y_title="", height=8, width=18):
    chart = LineChart()
    chart.title = title
    chart.style = 2
    chart.y_axis.title = y_title
    chart.height = height
    chart.width = width
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    for s in chart.series:
        s.smooth = False
    ws.add_chart(chart, anchor)


def add_bar_chart(ws, title, data_ref, cats_ref, anchor, y_title="", height=8, width=18):
    chart = BarChart()
    chart.type = "col"
    chart.title = title
    chart.style = 10
    chart.y_axis.title = y_title
    chart.height = height
    chart.width = width
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws.add_chart(chart, anchor)


def main():
    monthly = pd.read_csv(f"{PROC}/monthly_financials.csv", parse_dates=["month"])
    pnl = pd.read_csv(f"{PROC}/pnl_monthly.csv", parse_dates=["month"])
    ue_city = pd.read_csv(f"{PROC}/unit_economics_by_city.csv")
    ue_segment = pd.read_csv(f"{PROC}/unit_economics_by_segment.csv")
    ue_channel = pd.read_csv(f"{PROC}/unit_economics_by_channel.csv")
    cac_ltv = pd.read_csv(f"{PROC}/cac_ltv_by_channel.csv")
    drivers = pd.read_csv(f"{PROC}/margin_driver_decomposition.csv")
    quadrants = pd.read_csv(f"{PROC}/restaurant_quadrants.csv")
    rest_summary = pd.read_csv(f"{PROC}/restaurant_summary.csv")
    rid_summary = pd.read_csv(f"{PROC}/rider_summary.csv")
    cust_enriched = pd.read_csv(f"{PROC}/customers_enriched.csv")
    cohorts = pd.read_csv(f"{PROC}/cohorts.csv", parse_dates=["cohort_month"])
    forecast = pd.read_csv(f"{PROC}/forecast.csv", parse_dates=["month"])
    alloc = pd.read_csv(f"{PROC}/optimization_allocation.csv")
    with open(f"{PROC}/optimization_summary.json") as f:
        opt_summary = json.load(f)
    with open(f"{PROC}/strategic_plan.json") as f:
        plan = json.load(f)
    with open(f"{ROOT}/outputs/reports/key_findings.json") as f:
        findings = json.load(f)

    wb = Workbook()
    wb.remove(wb.active)

    # ---------------- Executive Summary ----------------
    ws = wb.create_sheet("Executive_Summary")
    sheet_title(ws, "NOVAFOOD — Executive Summary", "Business Intelligence & Financial Planning Case Study | Synthetic data", width=8)
    kpis = [
        ("GMV (trailing month)", f"${monthly.gmv.iloc[-1]:,.0f}"),
        ("Delivered Orders (trailing month)", f"{monthly.delivered_orders.iloc[-1]:,.0f}"),
        ("Contribution Margin % (trailing 3mo avg)", f"{monthly.contribution_margin_pct.tail(3).mean()*100:.1f}%"),
        ("Annualized Order Growth Rate", f"{findings['annualized_order_growth_rate_pct']:.1f}%"),
        ("Active Customers (trailing month)", f"{monthly.active_customers.iloc[-1]:,.0f}"),
        ("On-Time Delivery Rate (trailing month)", f"{monthly.on_time_rate.iloc[-1]*100:.1f}%"),
        ("Highest-Margin City", f"{findings['highest_margin_city']['city']} ({findings['highest_margin_city']['contribution_margin_pct']:.1f}%)"),
        ("Weakest-Margin City", f"{findings['weakest_margin_city']['city']} ({findings['weakest_margin_city']['contribution_margin_pct']:.1f}%)"),
    ]
    r = 4
    for label, val in kpis:
        ws.cell(row=r, column=1, value=label).font = LABEL_FONT
        vcell = ws.cell(row=r, column=4, value=val)
        vcell.font = Font(name=FONT, size=12, bold=True, color=TEAL)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Central management question").font = LABEL_FONT
    r += 1
    ws.cell(row=r, column=1, value="Can NOVAFOOD grow orders while improving contribution profit?").font = Font(name=FONT, italic=True, size=11)
    r += 2
    ws.cell(row=r, column=1, value="Strategic plan verdict (25% order growth, +2pp margin target)").font = LABEL_FONT
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.cell(row=r, column=1, value=plan["verdict"]).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 45
    r += 2
    ws.cell(row=r, column=1, value="Key findings (calculated from the dataset)").font = LABEL_FONT
    r += 1
    finding_lines = [
        f"- {findings['best_ltv_cac_channel']['channel']} is the strongest acquisition channel by LTV:CAC ({findings['best_ltv_cac_channel']['ltv_to_cac']:.2f}x).",
        f"- {findings['highest_value_segment']['segment']} customers generate the highest contribution profit per order (${findings['highest_value_segment']['contribution_profit_per_order']:.2f}).",
        f"- A promotion-budget re-allocation (this month's spend) could lift incremental contribution profit by ~{findings['optimization_incremental_profit_gain_pct']:.0f}% versus an even split across segments/cities.",
        f"- {findings['weakest_margin_city']['city']} runs a negative-to-thin contribution margin and is the priority market for operational review.",
    ]
    for line in finding_lines:
        ws.cell(row=r, column=1, value=line).alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        r += 1
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["D"].width = 26

    # ---------------- Financials ----------------
    ws = wb.create_sheet("Financials")
    sheet_title(ws, "Monthly P&L", "Contribution-margin marketplace P&L, computed from order-level ledger data", width=15)
    pnl_disp = pnl.copy()
    end_row = write_df(ws, pnl_disp, 4, currency_cols=[c for c in pnl_disp.columns if c not in ("month", "contribution_margin_pct")],
                        pct_cols=["contribution_margin_pct"])
    cats = Reference(ws, min_col=1, min_row=5, max_row=end_row - 1)
    data = Reference(ws, min_col=2, max_col=3, min_row=4, max_row=end_row - 1)
    add_line_chart(ws, "GMV vs Discounts by Month", data, cats, "A" + str(end_row + 2), y_title="$")
    data2 = Reference(ws, min_col=9, max_col=9, min_row=4, max_row=end_row - 1)
    add_line_chart(ws, "Contribution Profit by Month", data2, cats, "K" + str(end_row + 2), y_title="$")

    # ---------------- Unit Economics ----------------
    ws = wb.create_sheet("Unit_Economics")
    sheet_title(ws, "Unit Economics", "AOV, revenue/order, variable cost/order, and contribution profit/order by city, segment, and acquisition channel", width=12)
    r0 = 4
    ws.cell(row=r0, column=1, value="By City").font = LABEL_FONT
    end_row = write_df(ws, ue_city.drop(columns=["city_id"]), r0 + 1, pct_cols=["contribution_margin_pct"],
                        currency_cols=["aov", "revenue_per_order", "variable_cost_per_order", "contribution_profit_per_order"])
    r1 = end_row + 2
    ws.cell(row=r1, column=1, value="By Customer Segment").font = LABEL_FONT
    end_row2 = write_df(ws, ue_segment, r1 + 1, pct_cols=["contribution_margin_pct"],
                         currency_cols=["aov", "revenue_per_order", "variable_cost_per_order", "contribution_profit_per_order"])
    r2 = end_row2 + 2
    ws.cell(row=r2, column=1, value="By Acquisition Channel").font = LABEL_FONT
    write_df(ws, ue_channel, r2 + 1, pct_cols=["contribution_margin_pct"],
             currency_cols=["aov", "revenue_per_order", "variable_cost_per_order", "contribution_profit_per_order"])

    # ---------------- Customers ----------------
    ws = wb.create_sheet("Customers")
    sheet_title(ws, "Customer Analytics", "Segmentation (rule-based on realized behavior), CAC/LTV by channel", width=12)
    seg_counts = cust_enriched.customer_segment.value_counts().reset_index()
    seg_counts.columns = ["customer_segment", "customers"]
    seg_profit = cust_enriched.groupby("customer_segment").lifetime_profit.mean().round(2).reset_index()
    seg_profit.columns = ["customer_segment", "avg_lifetime_profit"]
    seg_table = seg_counts.merge(seg_profit, on="customer_segment")
    end_row = write_df(ws, seg_table, 4, currency_cols=["avg_lifetime_profit"])
    cats = Reference(ws, min_col=1, min_row=5, max_row=end_row - 1)
    data = Reference(ws, min_col=2, max_col=2, min_row=4, max_row=end_row - 1)
    add_bar_chart(ws, "Customers by Segment", data, cats, "A" + str(end_row + 2))

    r2 = end_row + 22
    ws.cell(row=r2, column=1, value="CAC / LTV by Acquisition Channel").font = LABEL_FONT
    end_row2 = write_df(ws, cac_ltv, r2 + 1, currency_cols=["cac", "ltv", "avg_lifetime_profit", "avg_lifetime_gmv"], pct_cols=["churn_rate"])

    # ---------------- Cohorts ----------------
    ws = wb.create_sheet("Cohorts")
    sheet_title(ws, "Monthly Retention Cohorts", "Share of each signup cohort still ordering N months later", width=14)
    pivot = cohorts.pivot_table(index="cohort_month", columns="month_number", values="retention_rate")
    pivot.columns = [f"M{int(c)}" for c in pivot.columns]
    pivot = pivot.reset_index()
    pivot["cohort_month"] = pivot.cohort_month.dt.strftime("%Y-%m")
    end_row = write_df(ws, pivot, 4, pct_cols=list(pivot.columns[1:]))
    rng = f"B5:{get_column_letter(pivot.shape[1])}{end_row-1}"
    ws.conditional_formatting.add(rng, ColorScaleRule(
        start_type="min", start_color="FFC7CE", mid_type="percentile", mid_value=50, mid_color="FFEB9C",
        end_type="max", end_color="C6EFCE"))

    # ---------------- Restaurants ----------------
    ws = wb.create_sheet("Restaurants")
    sheet_title(ws, "Restaurant Performance", "Revenue vs. margin quadrant classification", width=14)
    quad_counts = quadrants.quadrant.value_counts().reset_index()
    quad_counts.columns = ["quadrant", "restaurants"]
    end_row = write_df(ws, quad_counts, 4)
    cats = Reference(ws, min_col=1, min_row=5, max_row=end_row - 1)
    data = Reference(ws, min_col=2, max_col=2, min_row=4, max_row=end_row - 1)
    add_bar_chart(ws, "Restaurants by Quadrant", data, cats, "A" + str(end_row + 2))
    r2 = end_row + 22
    ws.cell(row=r2, column=1, value="Top 20 Restaurants by Contribution Profit").font = LABEL_FONT
    top20 = quadrants.sort_values("contribution_profit", ascending=False).head(20)[
        ["restaurant_name", "cuisine", "restaurant_tier", "city_id", "monthly_orders", "monthly_gmv",
         "contribution_profit", "margin_pct", "quadrant"]]
    write_df(ws, top20, r2 + 1, currency_cols=["monthly_gmv", "contribution_profit"], pct_cols=["margin_pct"])

    # ---------------- Operations ----------------
    ws = wb.create_sheet("Operations")
    sheet_title(ws, "Delivery Operations", "Rider-level performance and month-over-month service levels", width=12)
    ops_monthly = monthly[["month", "avg_delivery_minutes", "on_time_rate", "cancellation_rate", "avg_rating"]]
    end_row = write_df(ws, ops_monthly, 4, pct_cols=["on_time_rate", "cancellation_rate"])
    cats = Reference(ws, min_col=1, min_row=5, max_row=end_row - 1)
    data = Reference(ws, min_col=3, max_col=3, min_row=4, max_row=end_row - 1)
    add_line_chart(ws, "On-Time Delivery Rate by Month", data, cats, "A" + str(end_row + 2), y_title="%")
    r2 = end_row + 22
    ws.cell(row=r2, column=1, value="Top 15 Riders by Orders Completed").font = LABEL_FONT
    top_riders = rid_summary.sort_values("orders_completed", ascending=False).head(15)[
        ["rider_id", "vehicle_type", "orders_completed", "average_delivery_time", "average_distance", "on_time_rate", "cancellation_rate"]]
    write_df(ws, top_riders, r2 + 1, pct_cols=["on_time_rate", "cancellation_rate"])

    # ---------------- Forecast ----------------
    ws = wb.create_sheet("Forecast")
    sheet_title(ws, "12-Month Demand Forecast", "Rolling-origin-validated model per metric; base/upside/downside scenario bands", width=10)
    fc_orders = forecast[forecast.metric == "delivered_orders"][["month", "point_forecast", "lower_80", "upper_80"]]
    end_row = write_df(ws, fc_orders, 4, number_formats={"point_forecast": "#,##0", "lower_80": "#,##0", "upper_80": "#,##0"})
    cats = Reference(ws, min_col=1, min_row=5, max_row=end_row - 1)
    data = Reference(ws, min_col=2, max_col=4, min_row=4, max_row=end_row - 1)
    add_line_chart(ws, "Orders Forecast (12 Months, 80% Interval)", data, cats, "A" + str(end_row + 2), y_title="Orders")

    r2 = end_row + 22
    ws.cell(row=r2, column=1, value="Scenario Cases — Delivered Orders").font = LABEL_FONT
    scn = forecast[forecast.metric == "delivered_orders"][["month", "base_case", "upside_case", "downside_case"]]
    end_row2 = write_df(ws, scn, r2 + 1, number_formats={c: "#,##0" for c in ["base_case", "upside_case", "downside_case"]})
    cats2 = Reference(ws, min_col=1, min_row=r2 + 2, max_row=end_row2 - 1)
    data2 = Reference(ws, min_col=2, max_col=4, min_row=r2 + 1, max_row=end_row2 - 1)
    add_line_chart(ws, "Scenario Cases — Orders", data2, cats2, "A" + str(end_row2 + 2), y_title="Orders")

    # ---------------- Scenarios (assumptions + P&L impact) ----------------
    ws = wb.create_sheet("Scenarios")
    sheet_title(ws, "Scenario Planning", "Base / Upside / Downside assumptions and resulting monthly contribution profit (avg. of forecast horizon)", width=10)
    base_p = forecast[forecast.metric == "contribution_profit"].base_case.mean()
    up_p = forecast[forecast.metric == "contribution_profit"].upside_case.mean()
    down_p = forecast[forecast.metric == "contribution_profit"].downside_case.mean()
    scn_table = pd.DataFrame([
        {"scenario": "Downside", "order_growth_assumption": "-15% vs base", "avg_monthly_contribution_profit": down_p},
        {"scenario": "Base", "order_growth_assumption": "Model point forecast", "avg_monthly_contribution_profit": base_p},
        {"scenario": "Upside", "order_growth_assumption": "+12% vs base", "avg_monthly_contribution_profit": up_p},
    ])
    end_row = write_df(ws, scn_table, 4, currency_cols=["avg_monthly_contribution_profit"])
    cats = Reference(ws, min_col=1, min_row=5, max_row=end_row - 1)
    data = Reference(ws, min_col=3, max_col=3, min_row=4, max_row=end_row - 1)
    add_bar_chart(ws, "Avg Monthly Contribution Profit by Scenario", data, cats, "A" + str(end_row + 2))

    r2 = end_row + 22
    ws.cell(row=r2, column=1, value="Margin Driver Decomposition (first month vs. latest month)").font = LABEL_FONT
    write_df(ws, drivers, r2 + 1, number_formats={"approx_margin_pp_impact": "0.00"})

    # ---------------- Optimization ----------------
    ws = wb.create_sheet("Optimization")
    sheet_title(ws, "Promotion Budget Optimization", "SciPy trust-constr allocation across city x segment cells vs. an even-split baseline", width=12)
    summ = pd.DataFrame([
        {"metric": "Total monthly promo budget optimized", "value": opt_summary["total_budget"]},
        {"metric": "Baseline (even-split) incremental profit", "value": opt_summary["baseline_incremental_profit"]},
        {"metric": "Optimized incremental profit", "value": opt_summary["optimized_incremental_profit"]},
        {"metric": "Incremental profit gain", "value": opt_summary["incremental_profit_gain"]},
        {"metric": "Gain %", "value": opt_summary["gain_pct"] / 100},
    ])
    end_row = write_df(ws, summ, 4, number_formats={"value": "#,##0.00"})
    r2 = end_row + 2
    ws.cell(row=r2, column=1, value="Allocation by City x Segment").font = LABEL_FONT
    alloc_disp = alloc[["city_name", "customer_segment", "orders", "baseline_spend", "optimized_spend",
                         "baseline_incremental_profit", "optimized_incremental_profit"]].sort_values(
        "optimized_incremental_profit", ascending=False)
    write_df(ws, alloc_disp, r2 + 1, currency_cols=["baseline_spend", "optimized_spend", "baseline_incremental_profit", "optimized_incremental_profit"])

    # ---------------- Data Dictionary ----------------
    ws = wb.create_sheet("Data_Dictionary")
    sheet_title(ws, "Data Dictionary", "Core tables and fields in the underlying dataset", width=6)
    dd = pd.DataFrame([
        ("orders.csv", "order_id", "Unique order identifier"),
        ("orders.csv", "customer_total_paid", "food_value + delivery_fee + service_fee - discount_amount"),
        ("orders.csv", "platform_revenue", "restaurant_commission + service_fee + 35% of delivery_fee, less refund offset"),
        ("orders.csv", "contribution_profit", "platform_revenue - rider_cost - payment_processing_cost - 50% of refund_amount"),
        ("orders.csv", "order_status", "Delivered / Cancelled / Refunded"),
        ("customers.csv", "customer_segment", "Rule-based on recency/frequency/monetary + subscription status (see docs/data_dictionary.md)"),
        ("restaurants.csv", "commission_rate", "Platform take rate on food_value, by restaurant tier"),
        ("riders.csv", "reliability", "Latent on-time propensity used to simulate delivery delay"),
        ("monthly_financials.csv", "contribution_margin_pct", "contribution_profit / platform_revenue"),
    ], columns=["table", "field", "definition"])
    write_df(ws, dd, 4)

    # ---------------- KPI Definitions ----------------
    ws = wb.create_sheet("KPI_Definitions")
    sheet_title(ws, "KPI Definitions", "", width=4)
    kpi_defs = pd.DataFrame([
        ("GMV", "Sum of food_value + delivery_fee + service_fee across delivered/refunded orders"),
        ("Take Rate", "platform_revenue / GMV"),
        ("Contribution Margin %", "contribution_profit / platform_revenue"),
        ("AOV", "GMV / delivered orders"),
        ("LTV", "Average customer lifetime contribution profit, by acquisition channel"),
        ("LTV:CAC", "LTV divided by assumed channel CAC"),
        ("On-Time Rate", "Share of delivered orders with delivery_delay_minutes <= 5"),
        ("WAPE", "Weighted Absolute Percentage Error — sum(|error|) / sum(|actual|), used for forecast model selection"),
    ], columns=["kpi", "definition"])
    write_df(ws, kpi_defs, 4)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("Saved", OUT)


if __name__ == "__main__":
    main()
